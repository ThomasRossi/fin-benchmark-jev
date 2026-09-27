"""Build the reviewer-facing Excel workbook from cached responses."""

import statistics
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sklearn.metrics import f1_score

from bench import config
from bench.parse import INVALID, parse, score, set_f1

BOLD = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDE3EA")
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
GREY_TEXT = Font(color="808080")


def _pred(rec: dict | None, bench) -> str:
    """Re-parse from the stored raw reply so parser fixes apply to past runs without new API calls."""
    if rec is None:
        return "MISSING"
    return "ERROR" if rec.get("error") else parse(bench, rec.get("raw_text"), rec.get("api", "chat"))


def _applicable(model, bench) -> bool:
    return not (model.api == "decisions" and bench.kind == "number")


def _header(ws, row: int, values: list[str]) -> None:
    for col, v in enumerate(values, 1):
        c = ws.cell(row=row, column=col, value=v)
        c.font, c.fill, c.alignment = BOLD, HEADER_FILL, WRAP


def _pct(q: float, xs: list[float]) -> float | None:
    if not xs:
        return None
    xs = sorted(xs)
    return xs[min(len(xs) - 1, round(q * (len(xs) - 1)))]


def _f1(bench, gold: list[str], pred: list[str]) -> float | None:
    if bench.kind == "label":
        return f1_score(gold, pred, labels=sorted(set(gold)), average="macro", zero_division=0)
    if bench.kind == "idset":
        return statistics.mean(set_f1(p, g) for p, g in zip(pred, gold))
    return None  # numeric answers: accuracy only


def _metrics(bench, model, items, cache) -> dict:
    recs = [cache.get((model.key, bench.key, it["id"])) for it in items]
    ok = [r for r in recs if r and not r.get("error")]
    gold = [it["gold"] for it in items]
    pred = [_pred(r, bench) for r in recs]
    lat = [r["latency_ms"] for r in ok]
    out_toks = sum(r["completion_toks"] for r in ok)
    costs = [r["cost"] for r in ok if r.get("cost") is not None]
    reasoning = [r["reasoning_toks"] for r in ok if r.get("reasoning_toks") is not None]
    if not costs:
        cost_source = "missing: enter manually"
    elif model.price_in is not None:
        cost_source = "computed from list price"
    else:
        cost_source = "API usage.cost" + ("" if len(costs) == len(ok) else f" ({len(costs)}/{len(ok)} calls)")
    return {
        "n": len(items),
        "accuracy": sum(score(bench, p, g) for p, g in zip(pred, gold)) / len(items),
        "f1": _f1(bench, gold, pred),
        "invalid": sum(p == INVALID for p in pred) / len(items),
        "errors": sum(p in ("ERROR", "MISSING") for p in pred),
        "mean": statistics.mean(lat) if lat else None,
        "p50": statistics.median(lat) if lat else None,
        "p95": _pct(0.95, lat),
        "tok_s": out_toks / (sum(lat) / 1000) if lat else None,
        "reasoning": statistics.mean(reasoning) if reasoning else None,
        "cost": sum(costs) if costs else None,
        "total_s": sum(lat) / 1000 if lat else None,
        "cost_1k": sum(costs) / len(costs) * 1000 if costs else None,
        "cost_source": cost_source,
    }


def _summary(ws, models, benches, samples, cache, n: int) -> None:
    ws.title = "Summary"
    ws["A1"] = "Jev vs Qwen 3.8 27B (Cerebras): constrained-output financial benchmarks"
    ws["A1"].font = Font(bold=True, size=14)
    cols = ["Benchmark", "Model", "Role", "n", "Accuracy", "Macro-F1 (labels) /\nmean set-F1 (IDs)", "Invalid %", "Errors",
            "Latency mean (ms)", "Latency p50 (ms)", "Latency p95 (ms)", "Total response time (s)",
            "Output tok/s", "Avg reasoning toks", "API/computed cost ($)", "Cost source",
            "Manual cost ($)\n(fill if blank)", "Cost used ($)", "Cost per 1k items ($)"]
    _header(ws, 3, cols)
    row = 4
    for b in benches:
        for m in models:
            if not _applicable(m, b):
                vals = [b.sheet, m.key, "primary" if m.primary else "secondary",
                        "N/A: decisions model can't output a number (TypeSafe docs: keep arithmetic in code)"]
                for col, v in enumerate(vals, 1):
                    ws.cell(row=row, column=col, value=v).font = GREY_TEXT
                row += 1
                continue
            k = _metrics(b, m, samples[b.key], cache)
            # Cost used = manual entry if given, else API/computed; per-1k derived from it by formula.
            used = f'=IF(Q{row}<>"",Q{row},IF(O{row}<>"",O{row},""))'
            per_1k = f'=IF(R{row}<>"",R{row}/D{row}*1000,"")'
            vals = [b.sheet, m.key, "primary" if m.primary else "secondary", k["n"], k["accuracy"], k["f1"], k["invalid"], k["errors"],
                    k["mean"], k["p50"], k["p95"], k["total_s"], k["tok_s"], k["reasoning"],
                    k["cost"], k["cost_source"], None, used, per_1k]
            for col, v in enumerate(vals, 1):
                c = ws.cell(row=row, column=col, value=v)
                if not m.primary:
                    c.font = GREY_TEXT
            for col in (5, 6, 7):
                ws.cell(row=row, column=col).number_format = "0.0%"
            for col in (9, 10, 11, 13, 14):
                ws.cell(row=row, column=col).number_format = "0"
            ws.cell(row=row, column=12).number_format = "0.0"
            for col in (15, 17, 18, 19):
                ws.cell(row=row, column=col).number_format = "0.0000"
            ws.cell(row=row, column=17).fill = PatternFill("solid", fgColor="FFF2CC")
            row += 1

    row += 1
    ws.cell(row=row, column=1, value="Model that actually answered (as reported in each response)").font = BOLD
    answered = Counter((r["model"], r.get("routed_model")) for r in cache.values() if not r.get("error"))
    for (key, name), count in sorted(answered.items(), key=lambda x: (x[0][0], -x[1])):
        row += 1
        ws.cell(row=row, column=1, value=key)
        ws.cell(row=row, column=2, value=name)
        ws.cell(row=row, column=3, value=count)

    row += 2
    ws.cell(row=row, column=1, value="Run metadata").font = BOLD
    meta = [("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
            ("Items per benchmark", n),
            ("Sampling", f"seed {config.SEED}. Labels: balanced round-robin per label. FinQA: random numeric-answer items. "
                         f"Multi-headline: {config.GROUP_SIZE} tweets per item from the Twitter train+validation splits, 1-5 on "
                         f"the target topic (no target tweet repeats until its topic is exhausted), topics rotated; "
                         f"filler never comes from related topics (config.NEIGHBOR_TOPICS).")]
    meta += [(f"Model: {m.key}", f"{m.label} | id={m.model_id} | params={m.extra_body}") for m in models]
    meta += [(f"Dataset: {b.sheet}", f"{b.hf_path}@{b.hf_revision[:10]} split={b.split}") for b in benches]
    meta += [("Response time", "wall-clock per request measured by the harness (send -> full response, "
                               "non-streaming, includes network + queueing + reasoning). Final successful attempt only.")]
    meta += [("Scoring", "Labels: parsed label == gold. FinQA: within 1% of gold, also accepting a x100 "
                         "percent/ratio mismatch (gold stores 14.46% as 0.1446). Multi-headline: exact ID set "
                         "(set-F1 shown as partial credit). INVALID = reply not parseable; ERROR = API call failed "
                         "after retries. Both count as wrong.")]
    for k, v in meta:
        row += 1
        ws.cell(row=row, column=1, value=k)
        ws.cell(row=row, column=2, value=str(v))
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 22
    for col in range(3, len(cols) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.column_dimensions["P"].width = 24
    ws.row_dimensions[3].height = 45
    ws.freeze_panes = "A4"


def _bench_sheet(wb, lists_ws, list_col: int, bench, models, items, cache) -> None:
    ws = wb.create_sheet(bench.sheet)
    cols = ["ID", "Input text", "Gold answer"]
    for m in models:
        tag = m.key if m.primary else f"{m.key} (secondary)"
        cols += [f"{tag}\noutput", f"{tag}\nparsed", f"{tag}\nscore", f"{tag}\nlatency ms"]
    primary_names = " & ".join(m.key for m in models if m.primary)
    human = "Human label" if bench.kind == "label" else "Human answer"
    router = next((m for m in models if m.key == "jev_router"), None)
    cols += (["Jev Router chose"] if router else []) + [f"Agree? ({primary_names})", human, "Human notes"]
    _header(ws, 1, cols)

    for r, it in enumerate(items, 2):
        preds = []
        vals = [it["id"], it["text"], it["gold"]]
        for m in models:
            if not _applicable(m, bench):
                vals += ["N/A", "N/A", None, None]
                continue
            rec = cache.get((m.key, bench.key, it["id"]))
            p = _pred(rec, bench)
            if m.primary:
                preds.append(p)
            raw = (rec.get("error") or rec.get("raw_text")) if rec else None
            vals += [raw, p, score(bench, p, it["gold"]), rec.get("latency_ms") if rec else None]
        if len(preds) < 2:
            agree = None  # fewer than two primary models answer this benchmark
        elif bench.kind == "number" and preds[0] not in (INVALID, "ERROR", "MISSING"):
            agree = "yes" if all(score(bench, p, preds[0]) for p in preds) else "no"  # same number within tolerance
        else:
            agree = "yes" if len(set(preds)) == 1 else "no"
        if router:
            rr = cache.get((router.key, bench.key, it["id"]))
            vals += [rr.get("routed_model") if rr else None]
        vals += [agree, None, None]
        for c, v in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=v).alignment = WRAP if c in (2, 4) else TOP
        for i, m in enumerate(models):
            if not m.primary:
                for c in range(4 + 4 * i, 8 + 4 * i):
                    ws.cell(row=r, column=c).font = GREY_TEXT

    last = len(items) + 1
    for i in range(len(models)):
        score_col = get_column_letter(4 + 4 * i + 2)
        rng = f"{score_col}2:{score_col}{last}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["1"], fill=GREEN))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["0"], fill=RED))

    if bench.kind == "label":
        # Dropdown of valid labels, sourced from the hidden Lists sheet (inline lists cap at 255 chars).
        for i, label in enumerate(bench.labels, 1):
            lists_ws.cell(row=i, column=list_col, value=label)
        lc = get_column_letter(list_col)
        dv = DataValidation(type="list", formula1=f"=Lists!${lc}$1:${lc}${len(bench.labels)}", allow_blank=True)
        human_col = get_column_letter(len(cols) - 1)
        dv.add(f"{human_col}2:{human_col}{last}")
        ws.add_data_validation(dv)

    widths = {1: 10, 2: 60, 3: 18}
    for c in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(c, 14)
    for i in range(len(models)):
        ws.column_dimensions[get_column_letter(4 + 4 * i)].width = 22
    ws.column_dimensions[get_column_letter(len(cols))].width = 40
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"


def _prompts(wb, benches) -> None:
    ws = wb.create_sheet("Prompts")
    _header(ws, 1, ["Benchmark", "System prompt", "User prompt (example with placeholder text)"])
    for r, b in enumerate(benches, 2):
        task = b.task.format(n=config.GROUP_SIZE, topic="<topic>") if b.kind == "idset" else b.task
        msgs = config.build_messages(b, {"text": "<input text>", "task": task})
        for c, v in enumerate([b.sheet, msgs[0]["content"], msgs[1]["content"]], 1):
            ws.cell(row=r, column=c, value=v).alignment = WRAP
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 80


def _errors(wb, cache) -> None:
    ws = wb.create_sheet("Errors")
    _header(ws, 1, ["Model", "Benchmark", "Item ID", "Error"])
    errs = [r for r in cache.values() if r.get("error")]
    for i, r in enumerate(errs, 2):
        for c, v in enumerate([r["model"], r["bench"], r["item_id"], r["error"]], 1):
            ws.cell(row=i, column=c, value=v)
    ws.column_dimensions["D"].width = 100


def build(models, benches, samples, cache, out_dir: Path, n: int, run_name: str) -> Path:
    # Keep only records for this run's models x benchmarks x sampled items (the raw log may hold stale ones).
    wanted = {(m.key, b.key, it["id"]) for b in benches for it in samples[b.key] for m in models}
    cache = {k: v for k, v in cache.items() if k in wanted}
    wb = Workbook()
    _summary(wb.active, models, benches, samples, cache, n)
    lists_ws = wb.create_sheet("Lists")
    for i, b in enumerate(benches, 1):
        _bench_sheet(wb, lists_ws, i, b, models, samples[b.key], cache)
    _prompts(wb, benches)
    _errors(wb, cache)
    wb.move_sheet("Lists", offset=len(wb.sheetnames))
    lists_ws.sheet_state = "hidden"
    # Timestamped and never overwritten: each build is a new file.
    stem = f"jev_vs_qwen_{run_name}_{datetime.now():%Y-%m-%d_%H%M%S}"
    out = out_dir / f"{stem}.xlsx"
    i = 2
    while out.exists():
        out = out_dir / f"{stem}_{i}.xlsx"
        i += 1
    wb.save(out)
    return out
